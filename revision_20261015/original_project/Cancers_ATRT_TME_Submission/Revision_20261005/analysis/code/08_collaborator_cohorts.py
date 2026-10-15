#!/usr/bin/env python3
"""Export frozen TME cohort memberships; do not rerun or change annotation."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd

BASE = Path(__file__).resolve().parents[1]
ROOT = BASE.parents[2]
OUT = BASE/'collaborator_handoff_20261006'
XLSX = BASE/'outputs/01a10a41-df64-7c83-abc5-c45b6e1af12c/TME_collaborator_cohorts_20261006.xlsx'
SOURCE = BASE/'tables/TME_revised_annotations.csv'
FLAGS = [
    'conservative_astrocytes', 'conservative_microglia', 'conservative_macrophage_like',
    'broader_astrocytes', 'broader_microglia_myeloid', 'broader_macrophage_BAM_like',
    'BAM_like_reference', 'macrophage_high_confidence_optional',
]
EXPECTED = [23, 15, 16, 47, 288, 83, 67, 9]

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()

def prepare():
    OUT.mkdir(exist_ok=True)
    (OUT/'cohorts').mkdir(exist_ok=True)
    source = pd.read_csv(SOURCE)
    original = pd.read_csv(ROOT/'NCBIGEO/snRNA/GSE283839_ATRT_RNA_V3_metadata.tsv', sep='\t')
    assert source.Cell.is_unique and len(source) == 516
    assert source.Final_Annotation_Focus.eq('TME').all()
    lookup = original.set_index('Cell').loc[source.Cell]
    for col in ['ID', 'subtype', 'Final_Annotation', 'Final_Annotation_Focus']:
        assert source[col].tolist() == lookup[col].tolist(), col
    counts_file=ROOT/'NCBIGEO/snRNA/GSE283839_ATRT_RNA_V3_counts_raw.tsv'
    with counts_file.open() as handle:
        raw_header=handle.readline().rstrip('\n\r')
    raw_ids=raw_header.split('\t')
    normalized=[re.sub(r'^([ACGT]+)\.([0-9]+)', r'\1-\2', x) for x in raw_ids]
    assert len(raw_ids)==17564 and len(set(normalized))==17564
    barcode_map=dict(zip(normalized,raw_ids))
    assert set(source.Cell).issubset(barcode_map)

    d=pd.DataFrame({
        'cell_id':source.Cell,
        'raw_counts_cell_id':source.Cell.map(barcode_map),
        'patient_id':source.ID,
        'patient_group':source.subtype.str.removeprefix('ATRT-'),
        'source_subtype':source.subtype,
        'published_annotation':source.Final_Annotation,
        'revised_annotation':source.tme_annotation_revised,
        'annotation_confidence':source.annotation_confidence,
        'added_QC_pass':source.tme_annotation_revised.ne('QC uncertain').astype(int),
        'conservative_population':'Not included',
    })
    assert set(d.patient_group)=={'MYC','SHH','TYR'}
    assert (d.patient_id.str.extract(r'^ATRT-(MYC|SHH|TYR)-')[0]==d.patient_group).all()
    d['conservative_astrocytes']=source.included_in_conservative_astrocytes.astype(int)
    d['conservative_microglia']=source.included_in_conservative_microglia.astype(int)
    d['conservative_macrophage_like']=source.tme_annotation_revised.eq('Macrophage-like').astype(int)
    d['broader_astrocytes']=source.Final_Annotation.eq('Astrocytes').astype(int)
    d['broader_microglia_myeloid']=source.Final_Annotation.eq('Microglia').astype(int)
    d['broader_macrophage_BAM_like']=source.tme_annotation_revised.isin(['Macrophage-like','Border-associated macrophage-like']).astype(int)
    d['BAM_like_reference']=source.tme_annotation_revised.eq('Border-associated macrophage-like').astype(int)
    d['macrophage_high_confidence_optional']=(d.conservative_macrophage_like.eq(1)&source.annotation_confidence.str.startswith('High')).astype(int)
    for flag, label in zip(FLAGS[:3], ['Astrocyte','Microglia','Macrophage-like; infiltration unproven']):
        d.loc[d[flag].eq(1),'conservative_population']=label
    d['compartment']='TME'
    assert d[FLAGS].sum().tolist()==EXPECTED
    assert d[FLAGS[:3]].sum(axis=1).le(1).all()
    for narrow,broad in zip(FLAGS[:3],FLAGS[3:6]):
        assert (d[narrow]<=d[broad]).all()
    assert d.loc[d[FLAGS[:3]].any(axis=1),'added_QC_pass'].eq(1).all()
    d=d.sort_values(['conservative_population','patient_group','patient_id','cell_id']).reset_index(drop=True)
    d.to_csv(OUT/'TME_cell_registry.csv',index=False)

    definitions=[
        ('conservative_astrocytes','Primary conservative','Original Astrocytes AND revised Astrocyte; unchanged prior conservative rule.','broader_astrocytes'),
        ('conservative_microglia','Primary conservative','Original Microglia AND revised Microglia; unchanged prior conservative rule.','broader_microglia_myeloid'),
        ('conservative_macrophage_like','Primary conservative','Revised Macrophage-like; excludes BAM-like, monocyte-like, unresolved, mixed and QC-flagged cells. Includes 9 high and 7 moderate operational-confidence calls. Infiltration is unproven.','broader_macrophage_BAM_like'),
        ('broader_astrocytes','Broader sensitivity','All 47 original Astrocytes, including 24 uncertain cells. Five newly supported astrocyte candidates are not added.','conservative_astrocytes'),
        ('broader_microglia_myeloid','Broader sensitivity','All 288 original Microglia, including revised macrophage/BAM-like, unresolved and other qualified profiles; includes 4 added-QC-flagged cells.','conservative_microglia'),
        ('broader_macrophage_BAM_like','Broader sensitivity','16 Macrophage-like plus 67 BAM-like cells. Tests inclusion of a distinct macrophage-like population; this is not a broader infiltrating-macrophage label.','conservative_macrophage_like'),
        ('BAM_like_reference','Separate reference','The 67 revised border-associated macrophage-like cells, kept separately identifiable. Expression resemblance does not establish anatomical location.','broader_macrophage_BAM_like'),
        ('macrophage_high_confidence_optional','Optional stricter subset','The 9 of 16 Macrophage-like cells with the pre-existing High (marker and reference signatures) confidence label. Marker/reference overlap limits independence; infiltration remains unproven.','conservative_macrophage_like'),
    ]
    defs=pd.DataFrame(definitions,columns=['cohort_id','role','definition','comparison_cohort'])
    defs['n_cells']=EXPECTED
    defs['n_patients']=[int(d.loc[d[f].eq(1),'patient_id'].nunique()) for f in FLAGS]
    defs.to_csv(OUT/'cohort_definitions.csv',index=False)
    member_cols=['cell_id','raw_counts_cell_id','patient_id','patient_group','source_subtype','published_annotation','revised_annotation','annotation_confidence','added_QC_pass','compartment']
    long=[]
    for f in FLAGS:
        part=d.loc[d[f].eq(1),member_cols].sort_values(['patient_group','patient_id','cell_id']).copy()
        part.to_csv(OUT/'cohorts'/f'{f}.csv',index=False)
        part.insert(0,'cohort_id',f);long.append(part)
    memberships=pd.concat(long,ignore_index=True)
    assert not memberships.duplicated(['cohort_id','cell_id']).any()
    memberships.to_csv(OUT/'TME_cohort_membership.csv',index=False)
    counts=d.groupby(['patient_group','patient_id'],sort=True)[FLAGS].sum().reset_index()
    counts.insert(2,'all_TME_nuclei',d.groupby(['patient_group','patient_id'],sort=True).size().to_numpy())
    counts.to_csv(OUT/'TME_patient_counts.csv',index=False)
    descriptions={
        'cell_id':'Exact Cell key in the source metadata and annotated TME object; preserve all suffixes.',
        'raw_counts_cell_id':'Exact corresponding column name in the deposited raw-count TSV; use for direct raw-matrix joins.',
        'patient_id':'Source ID; identifies the individual patient/sample used for donor-level analyses.',
        'patient_group':'Molecular subtype MYC, SHH or TYR, derived from source subtype; not an individual patient identifier.',
        'source_subtype':'Unmodified source subtype (ATRT-MYC, ATRT-SHH or ATRT-TYR).',
        'published_annotation':'Unmodified source Final_Annotation.',
        'revised_annotation':'Frozen TME refinement from the previous analysis; not rerun for this export.',
        'annotation_confidence':'Existing operational confidence, not a classification probability or proof of origin.',
        'added_QC_pass':'1 passes the added conservative QC safeguards; 0 was QC uncertain. Broad original cohorts retain such cells.',
        'conservative_population':'Disjoint primary conservative grouping; Not included means no membership in the three primary conservative sets.',
        'compartment':'TME for every row; no malignant/non-TME cells appear in this export.'}
    descriptions.update({row.cohort_id:'Membership: 1 included, 0 not included. '+row.definition for row in defs.itertuples()})
    dictionary=pd.DataFrame([{'column':col,'type':'integer_0_or_1' if col in FLAGS+['added_QC_pass'] else 'text','meaning':descriptions[col]} for col in d.columns])
    dictionary.to_csv(OUT/'data_dictionary.csv',index=False)
    (OUT/'README.txt').write_text('''GSE283839 TME collaborator cohorts | 6 October 2026

Purpose: frozen cell lists for conservative-versus-broader sensitivity analyses.
No annotation was rerun. Only 516 author-defined TME nuclei are included; the
original malignant/nonmalignant boundary and all non-TME labels remain unchanged.

PRIMARY COMPARISONS
Astrocytes: 23 conservative versus 47 original broader astrocytes.
Microglia: 15 conservative versus 288 original microglia/myeloid cells.
Macrophage-like: 16 conservative operational calls versus 83 macrophage/BAM-like
cells (16 + 67). No cell is identified here as a confirmed infiltrating macrophage.
The 67 BAM-like cells are exported separately; the broader macrophage comparison
tests their inclusion. An optional 9-cell high-confidence macrophage-like subset
is supplied separately; 7 of the 16 primary macrophage-like calls are Moderate.

FILES AND JOIN KEYS
TME_cell_registry.csv: one row per TME cell; numeric 1/0 membership flags.
TME_cohort_membership.csv: one row per cohort_id/cell_id membership, sorted by
cohort, patient group and patient ID. Repeated cells across cohorts are intentional.
cohorts/: eight ready-to-use cohort CSVs, each with patient and group labels.
TME_patient_counts.csv: cell counts for each cohort, with zero counts retained.
cohort_definitions.csv and data_dictionary.csv: exact definitions and column meanings.
TME_collaborator_cohorts_20261006.xlsx: readable workbook with the same cell registry,
formula-derived patient counts and cohort definitions.

cell_id matches metadata Cell and the annotated TME object. raw_counts_cell_id
matches the deposited raw-count TSV exactly. Only the initial numeric barcode
suffix differs between these formats (e.g. -1 versus .1); both exact IDs are
provided. Do not strip the appended suffixes or shorten either barcode.
patient_id is the original ID (for example ATRT-SHH-4). patient_group is the
molecular group MYC, SHH or TYR; source_subtype retains ATRT-MYC/SHH/TYR.

ANALYSIS USE
Use each cohort separately for the intended conservative/broader rerun. The three
primary conservative sets are disjoint; each is nested within its paired broader
set. The broader microglia/myeloid set contains all 83 broader macrophage/BAM-like
cells, so those broader sets MUST NOT be treated as independent cell types in a
single analysis or concatenated without handling overlap. Use the disjoint
conservative_population field for a three-population analysis.

The broader 288-cell set reproduces the published membership and includes four
cells flagged by the added QC safeguards. added_QC_pass allows a separately
declared QC sensitivity; do not silently change the broader set. For patient-group
comparisons, retain patient_id as the replicate and report donor-specific cell
counts. Do not treat individual nuclei as independent patient replicates.

Confidence denotes operational transcriptomic evidence. Expression profiles do
not prove infiltration, embryologic origin, anatomical location or function.
Reference signatures overlap the curated markers. No donor eligibility threshold
was imposed on these lists; apply the chosen downstream eligibility rules after
subsetting and report which patients remain.

Provenance: GSE283839 deposited metadata and raw-count column names, plus the
TME_revised_annotations.csv generated in the 5 October 2026 annotation analysis.
manifest.json records source and output hashes and validation results. The
cohort definitions preserve that analysis and do not change confidence labels.
''')
    inputs={'registry':{'columns':list(d.columns),'rows':d.values.tolist()},
            'patients':{'columns':list(counts.columns),'rows':counts.values.tolist()},
            'definitions':{'columns':list(defs.columns),'rows':defs.values.tolist()},
            'dictionary':dictionary.values.tolist()}
    (BASE/'cache/cohort_workbook_inputs.json').write_text(json.dumps(inputs,indent=2))
    manifest={'date':'2026-10-06','source_annotation_sha256':sha(SOURCE),
        'source_metadata_sha256':sha(ROOT/'NCBIGEO/snRNA/GSE283839_ATRT_RNA_V3_metadata.tsv'),
        'raw_count_header_sha256':hashlib.sha256(raw_header.encode()).hexdigest(),
        'annotation_rerun':False,'TME_cells':516,'patients':int(d.patient_id.nunique()),
        'cohort_counts':dict(zip(FLAGS,EXPECTED)),
        'checks':{'unique_cell_ids':True,'unique_raw_matrix_ids':bool(d.raw_counts_cell_id.is_unique),
        'patient_and_subtype_match_source':True,'only_author_TME_cells':True,
        'conservative_sets_disjoint':True,'conservative_nested_in_broader':True,
        'conservative_all_added_QC_pass':True,'broader_myeloid_macrophage_overlap':83},
        'files':{str(p.relative_to(OUT)):{'bytes':p.stat().st_size,'sha256':sha(p)}
                 for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='manifest.json'}}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2))
    print(defs[['cohort_id','n_cells','n_patients']].to_string(index=False))

def finalize():
    from openpyxl import load_workbook
    inputs=json.loads((BASE/'cache/cohort_workbook_inputs.json').read_text())
    wb=load_workbook(XLSX,read_only=True,data_only=True)
    assert list(wb['Cell registry'].values)==[tuple(inputs['registry']['columns'])]+[tuple(r) for r in inputs['registry']['rows']]
    actual=list(wb['Patient counts'].iter_rows(min_row=5,max_row=16,min_col=1,max_col=11,values_only=True))
    assert actual==[tuple(r) for r in inputs['patients']['rows']],(actual,inputs['patients']['rows'])
    assert list(wb['Patient counts'].iter_rows(min_row=17,max_row=17,min_col=3,max_col=11,values_only=True))[0]==tuple([516]+EXPECTED)
    actual_defs=list(wb['Cohort definitions'].iter_rows(min_row=2,max_row=9,min_col=3,max_col=4,values_only=True))
    assert actual_defs==[(r[4],r[5]) for r in inputs['definitions']['rows']]
    registry=pd.read_csv(OUT/'TME_cell_registry.csv').set_index('cell_id')
    long=pd.read_csv(OUT/'TME_cohort_membership.csv')
    for flag,n_expected in zip(FLAGS,EXPECTED):
        cohort=pd.read_csv(OUT/'cohorts'/f'{flag}.csv').set_index('cell_id')
        expected=registry.loc[registry[flag].eq(1)]
        assert len(cohort)==n_expected and cohort.index.is_unique
        assert set(cohort.index)==set(expected.index)==set(long.loc[long.cohort_id.eq(flag),'cell_id'])
        pd.testing.assert_frame_equal(cohort,expected.loc[cohort.index,cohort.columns],check_dtype=False)
    for ws in wb:
        assert not any(cell.data_type=='e' for row in ws for cell in row)
    source=pd.read_csv(SOURCE)
    assert sha(SOURCE)==json.loads((OUT/'manifest.json').read_text())['source_annotation_sha256']
    assert len(source)==516
    manifest=json.loads((OUT/'manifest.json').read_text())
    manifest['checks']['workbook_values_and_cached_patient_counts_match_CSV']=True
    manifest['checks']['spreadsheet_formula_errors']=0
    manifest['checks']['all_individual_and_long_cohort_files_match_registry']=True
    manifest['files'][XLSX.name]={'bytes':XLSX.stat().st_size,'sha256':sha(XLSX)}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2))
    archive=BASE/'TME_collaborator_handoff_20261006.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file(): z.write(p,'TME_collaborator_handoff_20261006/'+str(p.relative_to(OUT)))
        z.write(XLSX,'TME_collaborator_handoff_20261006/'+XLSX.name)
    with zipfile.ZipFile(archive) as z: assert z.testzip() is None
    print('Verified workbook, source preservation and ZIP:',archive)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--finalize',action='store_true')
    if parser.parse_args().finalize: finalize()
    else: prepare()
